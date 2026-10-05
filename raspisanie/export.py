"""Выгрузка: Excel (в духе бумажного расписания), iCalendar (.ics) для групп и преподавателей."""
from __future__ import annotations

import os
import re
import zlib
from datetime import date, datetime, timedelta

from .solver import parse_ru_date

SLOT_TIMES = [("09:00", "10:30"), ("10:40", "12:10"), ("12:40", "14:10"), ("14:20", "15:50"),
              ("16:00", "17:30"), ("17:40", "19:10"), ("19:20", "20:50")]
# Неделя 1 — та, что начинается 31.08.2026 (проверено по датам в PDF: 09.09 стоит во 2-й неделе,
# 14.09 и 28.09 — в 1-й).
WEEK1_MONDAY = date(2026, 8, 31)


def week_of(d: date) -> int:
    return 1 if ((d - WEEK1_MONDAY).days // 7) % 2 == 0 else 2


def event_text(e: dict) -> str:
    parts = []
    for l in e["lessons"]:
        s = (l["elective"] + " " if l["elective"] else "") + l["subject"]
        if l["kind"]:
            s += f" ({l['kind']}{' ' + l['dates'] if l['dates'] else ''})"
        t = ", ".join(t["name"] + (f" ({','.join(map(str, t['subgroups']))} п/гр)" if t["subgroups"] else "")
                      for t in l["teachers"])
        if t:
            s += "\n" + t
        parts.append(s)
    txt = "\n".join(parts)
    if e["online"]:
        txt += "\nдистанционно (el.mpgu.su)"
    elif e["offsite"]:
        txt += "\n" + e["offsite"]
    elif e["rooms"]:
        txt += "\nауд. " + ", ".join(e["rooms"])
    return txt


# ---------------------------------------------------------------- Excel
def to_xlsx(data: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    by_file: dict = {}
    for g, info in data["groups"].items():
        by_file.setdefault(info.get("course"), []).append(g)
    for course in sorted(by_file, key=lambda c: c or 0):
        groups = by_file[course]
        ws = wb.create_sheet(f"{course} курс" if course else "прочие")
        ws.append(["День", "Время", "Нед."] + groups)
        for c in ws[1]:
            c.font, c.fill, c.alignment, c.border = Font(bold=True), head_fill, wrap, border
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 5
        for j in range(len(groups)):
            ws.column_dimensions[ws.cell(1, 4 + j).column_letter].width = 34
        row_of = {}
        r = 2
        for d, dname in enumerate(data["days"]):
            for s, slot in enumerate(data["slots"]):
                for w in (1, 2):
                    row_of[(d, s, w)] = r
                    ws.cell(r, 1, dname)
                    ws.cell(r, 2, slot)
                    ws.cell(r, 3, w)
                    r += 1
        for e in data["events"]:
            cols = sorted(4 + groups.index(g) for g in e["groups"] if g in groups)
            if not cols:
                continue
            r0 = row_of[(e["day"], e["slot"], min(e["weeks"]))]
            r1 = row_of[(e["day"], e["slot"], max(e["weeks"]))]
            c0, c1 = cols[0], cols[-1]
            cell = ws.cell(r0, c0)
            prev = cell.value
            cell.value = (prev + "\n---\n" if prev else "") + event_text(e)
            if (r0, c0) != (r1, c1) and cols == list(range(c0, c1 + 1)) and not prev:
                try:
                    ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
                except ValueError:
                    pass
        for row in ws.iter_rows(min_row=2, max_row=r - 1, max_col=3 + len(groups)):
            for c in row:
                c.alignment, c.border = wrap, border
        for rr in range(2, r):
            ws.row_dimensions[rr].height = 48
        ws.freeze_panes = "D2"
    wb.save(path)


# ---------------------------------------------------------------- iCalendar
def _dates_in(spec: str, start: date, end: date) -> tuple[list[date] | None, date, date, list[date]]:
    """Разбирает «12.09, 19.09», «по 14.10», «с 31.10», «кроме 09.09, 23.09»."""
    year = start.year

    def mk(dd: str) -> date:
        d_, m_ = map(int, dd.split("."))
        y = year if m_ >= start.month else year + 1
        return date(y, m_, d_)

    found = re.findall(r"\d{1,2}\.\d{2}", spec)
    if not spec:
        return None, start, end, []
    if "кроме" in spec:
        return None, start, end, [mk(x) for x in found]
    m = re.search(r"\bпо\s+(\d{1,2}\.\d{2})", spec)
    if m:
        end = min(end, mk(m.group(1)))
        rest = [x for x in found if x != m.group(1)]
        return ([mk(x) for x in rest] or None) if rest else None, start, end, []
    m = re.search(r"\bс\s+(\d{1,2}\.\d{2})", spec)
    if m:
        return None, max(start, mk(m.group(1))), end, []
    return [mk(x) for x in found] or None, start, end, []


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def to_ics(data: dict, events: list[dict], name: str) -> str:
    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//matfak-mpgu//raspisanie//RU",
             f"X-WR-CALNAME:{_esc(name)}", "X-WR-TIMEZONE:Europe/Moscow"]
    for e in events:
        if e.get("kind") == "selfstudy":
            continue
        periods = [data["groups"][g]["period"] for g in e["groups"] if data["groups"].get(g, {}).get("period")]
        if not periods:
            continue
        start = min(parse_ru_date(p[0]) for p in periods)
        end = max(parse_ru_date(p[1]) for p in periods)
        spec = " ".join(l["dates"] for l in e["lessons"] if l["dates"])
        explicit, start, end, exclude = _dates_in(spec, start, end)
        t0, t1 = SLOT_TIMES[e["slot"]]
        if explicit:
            dates = [d for d in explicit if d.weekday() == e["day"]] or explicit
        else:
            first = start + timedelta(days=(e["day"] - start.weekday()) % 7)
            while len(e["weeks"]) == 1 and week_of(first) != e["weeks"][0]:
                first += timedelta(days=7)
            step = 7 if len(e["weeks"]) == 2 else 14
            dates, d = [], first
            while d <= end:
                if d not in exclude:
                    dates.append(d)
                d += timedelta(days=step)
        title = " + ".join(f"{l['subject']}{' (' + l['kind'] + ')' if l['kind'] else ''}" for l in e["lessons"])
        loc = "el.mpgu.su" if e["online"] else (e["offsite"] or ("Краснопрудная ул., 14, ауд. " + ", ".join(e["rooms"]) if e["rooms"] else ""))
        for d in dates:
            uid = f"{e['id']}-{d.isoformat()}-{zlib.crc32(name.encode()) % 10**8}@matfak-mpgu"
            lines += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}",
                      f"DTSTART;TZID=Europe/Moscow:{d:%Y%m%d}T{t0.replace(':', '')}00",
                      f"DTEND;TZID=Europe/Moscow:{d:%Y%m%d}T{t1.replace(':', '')}00",
                      f"SUMMARY:{_esc(title)}", f"LOCATION:{_esc(loc)}",
                      f"DESCRIPTION:{_esc(event_text(e))}", "END:VEVENT"]
    lines.append("END:VCALENDAR")
    return "\r\n".join(lines) + "\r\n"


def export_ics_dir(data: dict, out_dir: str) -> int:
    os.makedirs(out_dir, exist_ok=True)
    n = 0
    for g in data["groups"]:
        evs = [e for e in data["events"] if g in e["groups"]]
        with open(os.path.join(out_dir, f"{g}.ics"), "w", encoding="utf-8", newline="") as f:
            f.write(to_ics(data, evs, f"Расписание {g}"))
        n += 1
    for t in data["teachers"]:
        evs = [e for e in data["events"] if t in e["teachers"]]
        fn = re.sub(r"[^\wА-Яа-яЁё.-]+", "_", t)
        with open(os.path.join(out_dir, f"преп_{fn}.ics"), "w", encoding="utf-8", newline="") as f:
            f.write(to_ics(data, evs, f"Расписание: {t}"))
        n += 1
    return n
